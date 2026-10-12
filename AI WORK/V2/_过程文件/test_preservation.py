"""Positive preservation tests and deliberate forbidden-change detection tests."""
from pathlib import Path
import io, zipfile, json, sys, importlib.util
from copy import deepcopy
from lxml import etree as E
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('pure_style',ROOT/'apply_style.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
SRC=ROOT.parent/'V1'/'交付文件'/'产业园投资测算.xlsx'
V2=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'_过程文件'/'样式试稿'/m.OUTPUT_NAME
before,after=SRC.read_bytes(),V2.read_bytes()
results={'input_sha256':m.sha(before),'output_sha256':m.sha(after),'negative_tests':[]}

def mutate(data,part,fn):
    zparts=m.package(data)
    root=m.xml(zparts[part]);fn(root)
    replacement=E.tostring(root,encoding='utf-8',xml_declaration=True)
    stream=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi,zipfile.ZipFile(stream,'w') as zo:
        zo.comment=zi.comment
        for item in zi.infolist():zo.writestr(deepcopy(item),replacement if item.filename==part else zi.read(item))
    return stream.getvalue()

ns=m.NS
fixtures=[
 ('参数数值','xl/worksheets/sheet2.xml',lambda r:setattr(r.find("s:sheetData/s:row/s:c[@r='B10']/s:v",ns),'text','88888')),
 ('公式字符','xl/worksheets/sheet8.xml',lambda r:setattr(r.find("s:sheetData/s:row/s:c[@r='D26']/s:f",ns),'text','1+1')),
 ('公式缓存','xl/worksheets/sheet8.xml',lambda r:setattr(r.find("s:sheetData/s:row/s:c[@r='D26']/s:v",ns),'text','123456')),
 ('数字格式','xl/styles.xml',lambda r:r.find('s:cellXfs',ns)[0].set('numFmtId','2')),
 ('冻结窗格','xl/worksheets/sheet2.xml',lambda r:r.find('s:sheetViews/s:sheetView/s:pane',ns).set('ySplit','8')),
 ('隐藏行','xl/worksheets/sheet2.xml',lambda r:r.find('s:sheetData/s:row',ns).set('hidden','1')),
 ('合并范围','xl/worksheets/sheet2.xml',lambda r:r.find('s:mergeCells/s:mergeCell',ns).set('ref','A5:H5')),
 ('数据验证范围','xl/worksheets/sheet2.xml',lambda r:r.find('s:dataValidations/s:dataValidation',ns).set('sqref','B99')),
 ('条件格式优先级','xl/worksheets/sheet10.xml',lambda r:r.find('s:conditionalFormatting/s:cfRule',ns).set('priority','999')),
 ('计算设置','xl/workbook.xml',lambda r:r.find('s:calcPr',ns).set('calcMode','manual')),
 ('图表系列引用','xl/charts/chart1.xml',lambda r:setattr(r.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}f'),'text',"'错误引用'!$A$1")),
 ('图表位置','xl/drawings/drawing1.xml',lambda r:setattr(r.find('.//{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}row'),'text','99')),
 ('命名范围','xl/workbook.xml',lambda r:setattr(r.find('s:definedNames/s:definedName',ns),'text',"'错误引用'!$A$1")),
 ('对齐缩进','xl/styles.xml',lambda r:r.find('s:cellXfs',ns)[1].find('s:alignment',ns).set('indent','1')),
]
for label,part,fn in fixtures:
    candidate=mutate(after,part,fn)
    report=m.verify_packages(before,candidate)
    assert not report['passed'],f'Forbidden change was not detected: {label}'
    results['negative_tests'].append({'mutation':label,'correctly_rejected':True,'detected_by':[d['check'] for d in report['differences']]})

# Independent parser reading only: no openpyxl writing/saving.
wa=load_workbook(SRC,data_only=False,read_only=False)
wb=load_workbook(V2,data_only=False,read_only=False)
va=load_workbook(SRC,data_only=True,read_only=False)
vb=load_workbook(V2,data_only=True,read_only=False)
diffs=[];n=0;fc=0
assert wa.sheetnames==wb.sheetnames
for sa,sb,sva,svb in zip(wa,wb,va,vb):
    for row in sa:
        for a in row:
            b=sb[a.coordinate]
            if a.value is not None:
                n+=1
                if (a.value,a.data_type,a.number_format)!=(b.value,b.data_type,b.number_format):diffs.append([sa.title,a.coordinate,'content'])
                if a.data_type=='f':
                    fc+=1
                    av,bv=sva[a.coordinate],svb[a.coordinate]
                    if (av.value,av.data_type)!=(bv.value,bv.data_type):diffs.append([sa.title,a.coordinate,'cache'])
    assert str(sa.merged_cells)==str(sb.merged_cells)
    assert sa.freeze_panes==sb.freeze_panes
    assert sa.sheet_state==sb.sheet_state
    assert sa.print_area==sb.print_area
    assert sa.print_title_rows==sb.print_title_rows and sa.print_title_cols==sb.print_title_cols
assert not diffs,diffs[:10]
results['independent_openpyxl']={'version':'3.1.5','checked_nonempty_cells':n,'checked_formulas_and_caches':fc,'differences':len(diffs),'read_only_operations':True}

# User-modified parameter survives; stale original caches are also intentionally
# retained, demonstrating that this script does not initialize or recalculate.
def changed_parameter(r):
    r.find("s:sheetData/s:row/s:c[@r='B10']/s:v",ns).text='125000'
parameter_fixture=mutate(before,'xl/worksheets/sheet2.xml',changed_parameter)
fixture_dir=ROOT/'_过程文件'/'参数保留测试'
fixture_dir.mkdir(exist_ok=True)
fixture=fixture_dir/'用户已修改参数.xlsx'
m.write_new(fixture,parameter_fixture) if not fixture.exists() else None
fixture_hash=m.sha(fixture.read_bytes())
out,rep=m.apply_style(fixture,fixture_dir)
assert m.verify_packages(fixture.read_bytes(),out.read_bytes())['passed']
assert m.sha(fixture.read_bytes())==fixture_hash
value=m.xml(m.package(out.read_bytes())['xl/worksheets/sheet2.xml']).find("s:sheetData/s:row/s:c[@r='B10']/s:v",ns).text
assert value=='125000'
first_hash=m.sha(out.read_bytes())
second,rep2=m.apply_style(fixture,fixture_dir)
assert second!=out and m.sha(out.read_bytes())==first_hash
assert m.verify_packages(fixture.read_bytes(),second.read_bytes())['passed']
results['parameter_preservation']={'input_cell':'核心参数输入!B10','input_saved_value':125000,'output_value':125000,'input_sha256_unchanged':True,'all_caches_retained':True}
results['same_name_output_test']={'second_output_has_unique_name':True,'first_output_sha256_unchanged':True}
assert m.sha(SRC.read_bytes())==m.sha(before)
results['v1_unchanged']=True
(ROOT/'_过程文件'/'补充验证结果.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in results.items() if k!='negative_tests'},ensure_ascii=False))
print('Forbidden mutations correctly rejected:',len(results['negative_tests']))
