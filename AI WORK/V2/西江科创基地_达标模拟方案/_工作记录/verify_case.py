from pathlib import Path
import json, hashlib, zipfile, copy, math
import openpyxl
from lxml import etree as E
from patch_case import parts,mapping,N,S

H=Path(__file__).resolve().parent
m=json.loads((H/'case_manifest.json').read_text(encoding='utf-8'))
native=json.loads((H/'Excel实际重算及情景结果.json').read_text(encoding='utf-8-sig'))
src=Path(m['source']);out=Path(m['root'])/'西江科创基地_投资测算_达标模拟.xlsx'
a=parts(src);b=parts(out);am=mapping(a);bm=mapping(b)
w1=openpyxl.load_workbook(src,data_only=False);w2=openpyxl.load_workbook(out,data_only=False)
values=openpyxl.load_workbook(out,data_only=True)
calc=openpyxl.load_workbook(H/'Excel重算副本.xlsx',data_only=True)
errors=[];numeric_checks=0;string_checks=0;max_abs_error=0
def test(cond,msg):
 if not cond:errors.append(msg)
def close(x,y,label):
 global numeric_checks,string_checks,max_abs_error
 if isinstance(x,(int,float)):
  numeric_checks+=1
  if not isinstance(y,(int,float)) or not math.isfinite(y):errors.append(f'{label}: 非数值 {y!r}');return
  diff=abs(x-y);max_abs_error=max(max_abs_error,diff)
  if diff>max(.00001,abs(x)*.00000001):errors.append(f'{label}: {x} != {y}')
 elif isinstance(x,str):
  string_checks+=1
  if x!=y:errors.append(f'{label}: {x!r} != {y!r}')
for pred,actual in zip(m['predictions'],native['Cases']):
 test(pred['name']==actual['name'],'情景名称不对应')
 for scenario in ['base','stress']:
  for key,val in pred[scenario]['metrics'].items():
   if key in actual[scenario]:close(val,actual[scenario][key],pred['name']+'/'+scenario+'/'+key)
  for group,grids in actual['grids'].items():
   expected=pred[scenario][{'plan':'plan','investment_annual':'investment_annual','income':'income','costs':'costs','financing':'financing','cashflow':'cashflow'}[group]]
   for key,arr in grids[scenario].items():
    if key not in expected:continue
    for year,(x,y) in enumerate(zip(expected[key],arr)):
     if x is not None:close(x,y,f'{pred["name"]}/{scenario}/{group}/{key}/{year}')
formulas_original={(s.title,c.coordinate):c.value for s in w1 for row in s for c in row if c.data_type=='f'}
formulas_final={(s.title,c.coordinate):c.value for s in w2 for row in s for c in row if c.data_type=='f'}
test(formulas_original==formulas_final,'原公式字符串／地址／数量有变化')
test(w1.sheetnames==w2.sheetnames,'工作表顺序有变化')
cache_checks=0;unrelated_cells=0
for s in w1:
 t=w2[s.title]
 test(s.sheet_state==t.sheet_state,f'{s.title}可见状态')
 test(s.freeze_panes==t.freeze_panes,f'{s.title}冻结窗格')
 test(set(map(str,s.merged_cells))<=set(map(str,t.merged_cells)),f'{s.title}原合并范围')
 for row in s:
  for c in row:
   if c.value is None:continue
   d=t[c.coordinate]
   test(c.number_format==d.number_format,f'{s.title}!{c.coordinate}数字格式')
   test(c.style_id==d.style_id,f'{s.title}!{c.coordinate}样式')
   if c.coordinate not in m['changes'].get(s.title,{}) and c.data_type!='f':
    unrelated_cells+=1;test(c.value==d.value,f'{s.title}!{c.coordinate}非授权内容')
 for address,value in m['changes'].get(s.title,{}).items():test(t[address].value==value,f'{s.title}!{address}输入内容')
 for tag in ['dataValidations','conditionalFormatting','sheetProtection','autoFilter','drawing','printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks']:
  x=E.fromstring(a[am[s.title]]).findall(S+tag);y=E.fromstring(b[bm[s.title]]).findall(S+tag)
  test([E.tostring(n,method='c14n') for n in x]==[E.tostring(n,method='c14n') for n in y],f'{s.title}/{tag}变化')
 for row in values[s.title]:
  for c in row:
   if c.data_type=='e':errors.append(f'Excel错误 {s.title}!{c.coordinate}: {c.value}')
 for row in t:
  for c in row:
   if c.data_type=='f':
    cache_checks+=1;test(values[s.title][c.coordinate].value==calc[s.title][c.coordinate].value,'缓存与Excel重算不一致'+s.title+c.coordinate)
for member in a:
 if member in am.values():continue
 if member.startswith('xl/charts/chart') and member.endswith('.xml'):
  def without_cache(data):
   t=E.fromstring(data)
   for node in t.xpath('//c:numCache|//c:strCache',namespaces=N):node.getparent().remove(node)
   return E.tostring(t,method='c14n')
  test(without_cache(a[member])==without_cache(b[member]),f'图表定义改变{member}')
 else:test(a[member]==b[member],f'非授权ZIP成员改变{member}')
base=native['Cases'][0]
test(all(r['base']=='通过' and r['stress']=='通过' for r in base['statuses']),'交付方案存在未通过指标')
test(all(c['value']=='通过' for c in base['checks']),'交付方案存在未通过校验')
test(base['decision']=='通过：当前模拟参数下具备投资条件','最终决策不通过')
for r in range(69,76):test(w1['核心参数输入'][f'B{r}'].value==w2['核心参数输入'][f'B{r}'].value,'阈值被调整')
test(w2['核心参数输入']['B64'].value==0,'处置开关未关闭')
test(hashlib.sha256(src.read_bytes()).hexdigest()==m['source_sha256'],'原V2哈希变化')
underfund=next(c for c in native['Cases'] if c['name']=='资本金额度不足')
test(underfund['base']['max_gap']>0 and underfund['decision']=='未通过','资本金不足未识别')
bad_price=next(c for c in native['Cases'] if c['name']=='售价跌至4000元')
test(bad_price['decision']=='未通过','售价不足未正确拒绝')
report=dict(passed=not errors,errors=errors,numeric_comparisons=numeric_checks,string_comparisons=string_checks,
            maximum_abs_error=max_abs_error,formula_count=len(formulas_final),formula_cache_checks=cache_checks,
            unchanged_unrelated_content_cells=unrelated_cells,source_sha256=m['source_sha256'],
            final_sha256=hashlib.sha256(out.read_bytes()).hexdigest(),ExcelVersion=native['ExcelVersion'],
            base=base['base'],stress=base['stress'],decision=base['decision'],statuses=base['statuses'],checks=base['checks'],
            scenarios=[dict(name=x['name'],decision=x['decision'],reasons=x['reasons'],base=x['base'],stress=x['stress']) for x in native['Cases']])
(Path(m['root'])/'方案验证结果.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:report[k] for k in ['passed','errors','numeric_comparisons','string_comparisons','maximum_abs_error','formula_count','formula_cache_checks','unchanged_unrelated_content_cells','decision']},ensure_ascii=False,indent=2))
if errors:raise SystemExit(1)
