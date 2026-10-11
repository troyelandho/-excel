"""Independent numerical oracle and actual Excel recalculation / mutation checks.

Run --prepare, then the emitted PowerShell script, then --check.
This validator is separate from workbook authoring and never supplies formula caches.
"""
from __future__ import annotations
import argparse
import json
import math
import re
import shutil
import sys
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries
import generate_excel as g

CASES = [
 ('基准',{}),
 ('售价下降10%',{'sale_price':5850}),
 ('出租率下降10个百分点',{'occ1':.40,'occ2':.60,'occ3':.75,'occ_stable':.80}),
 ('非土地成本增加10%',{'above_cost':2640,'under_cost':3300,'infra':4950,'pre':1980,'other_build':2420}),
 ('利率提高至6.5%',{'interest':.065}),
 ('资本金额度不足',{'equity_cap':20000}),
 ('期限12年',{'horizon':12}),
 ('期限30年',{'horizon':30}),
 ('期末处置开启',{'terminal_on':1}),
 ('无贷款',{'loan':0}),
 ('要求压力也通过',{'stress_required':1}),
 ('期限不覆盖贷款',{'horizon':8}),
 ('销售仅完成70%',{'sale_plan':[0,0,0,0,.2,.3,.2]+[0]*24}),
 ('出租推迟至第6年',{'rent_start':6}),
 ('无正现金流IRR无解',{'sale_price':0,'rent':0}),
 ('保本出租率超过100%',{'fixed_cost':2000}),
 ('模拟全部条件通过',{'sale_price':11000,'rent':75,'loan':15000,'equity_cap':60000,
                         'static_limit':20,'dynamic_limit':30}),
 ('基准通过但要求压力通过',{'sale_price':11000,'rent':75,'loan':15000,'equity_cap':60000,
                              'static_limit':20,'dynamic_limit':30,'stress_required':1}),
 ('贷款到期超出预留期',{'loan_term':40})]

def irr(cf):
    """Find a root by discount-factor polynomial, not Excel's iterative IRR routine."""
    if not (min(cf)<0<max(cf)):return '无法计算'
    cf=np.trim_zeros(np.array(cf,dtype=float),'b')
    roots=np.polynomial.polynomial.polyroots(cf)
    rates=[1/x.real-1 for x in roots if abs(x.imag)<1e-7 and x.real>0]
    if not rates:return '无法计算'
    return min(rates,key=lambda x:abs(x-.1))

def payback(cf, rate=0):
    cf=np.asarray(cf)/(1+rate)**np.arange(len(cf))
    if not any(cf<0):return '不适用'
    running=0.0
    for year,flow in enumerate(cf):
        previous=running;running+=flow
        if year>0 and previous<0 and running>=0:return year-1-previous/flow
    return '测算期内未回收'

def oracle(overrides=None, stress=False):
    p={**g.DEFAULTS,**(overrides or {})};n=int(p['horizon']);y=np.arange(31);on=y<=n
    build=np.array(p.get('build_plan',g.DEFAULT_BUILD_PLAN),dtype=float)
    sales_plan=np.array(p.get('sale_plan',g.DEFAULT_SALE_PLAN),dtype=float)
    build=build*((y>=1)&(y<=p['build_years'])&on)
    sales_plan=sales_plan*((y>=p['sale_start'])&on)
    land=p['land_area'];above=land*p['far'];operable=above*p['operable_share']
    sold_area=operable*p['sale_share'];held_area=operable*(1-p['sale_share'])
    factor=1+p['stress_cost'] if stress else 1
    land_cost=land*p['land_price']/10000
    above_cost=above*p['above_cost']/10000*factor
    under_cost=p['under_area']*p['under_cost']/10000*factor
    pre_cont=above_cost+under_cost+(p['infra']+p['pre']+p['other_build'])*factor
    nonland=pre_cont*(1+p['contingency'])
    draw=build*p['loan']
    first_draw=next((int(t) for t in y if draw[t]>1e-6),0)
    maturity=first_draw+p['loan_term']-1 if p['loan']>0 else 0
    count=max(1,maturity-p['repay_start']+1) if p['loan'] else 0
    principal=np.zeros(31);opening=np.zeros(31);closing=np.zeros(31)
    for t in range(n+1):
        opening[t]=closing[t-1] if t else 0
        if p['loan']>0 and p['repay_start']<=t<=maturity:
            principal[t]=min(opening[t]+draw[t],p['loan']/max(1,count))
        closing[t]=max(0,opening[t]+draw[t]-principal[t])
    interest=(opening+draw/2-principal/2).clip(0)*p['interest']*on
    cap_interest=interest*(y<=p['build_years'])
    op_interest=interest-cap_interest;service=interest+principal
    capitalized=sum(cap_interest);development=land_cost+nonland+capitalized
    sale_basis_project=(land_cost+nonland)*p['sale_share']
    sale_basis_equity=development*p['sale_share']
    dep_basis_project=nonland*(1-p['sale_share'])
    dep_basis_equity=(nonland+capitalized)*(1-p['sale_share'])
    held_land=land_cost*(1-p['sale_share'])
    held_invest=development*(1-p['sale_share'])+p['working_capital']
    age=(y-p['rent_start']+1).clip(0)*on
    ramp=[p['occ1'],p['occ2'],p['occ3']]
    occ=np.array([0 if not a else ramp[int(a)-1] if a<=3 else p['occ_stable'] for a in age])
    if stress:occ=np.maximum(0,occ-p['stress_occ']/100)*(age>0)
    occupied=held_area*occ
    new=np.maximum(0,np.diff(occupied,prepend=0))
    # Outside horizon all cash schedules are inactive, including new occupancy.
    new*=on
    price=p['rent']*(1+p['rent_growth'])**np.maximum(0,age-1)*(age>0)
    if stress:price*=1-p['stress_rent']
    free_loss=new*price*p['free_months']/10000
    rent=(occupied*price*12/10000-free_loss)*p['collection']
    other=rent*p['other_income'];held_income=rent+other
    sale_price=p['sale_price']*(1-p['stress_sale'] if stress else 1)
    sale_revenue=sold_area*sales_plan*sale_price/10000
    revenue=sale_revenue+held_income
    op_age=np.maximum(0,y-p['build_years'])
    fixed=p['fixed_cost']*(1+p['fixed_growth'])**np.maximum(0,op_age-1)*(op_age>0)*on
    maintenance=p['maintenance']*(op_age>0)*on
    variable=held_income*p['variable_cost'];op_tax=held_income*p['operating_tax']
    sale_expense=sale_revenue*p['sale_expense'];sale_tax=sale_revenue*p['sale_tax']
    cash_cost=fixed+maintenance+variable+op_tax+sale_expense+sale_tax
    sale_cost_project=sale_basis_project*sales_plan
    sale_cost_equity=sale_basis_equity*sales_plan
    dep_active=(op_age>0)&(op_age<=p['dep_years'])&on
    dep_project=dep_basis_project*(1-p['residual'])/p['dep_years']*dep_active
    dep_equity=dep_basis_equity*(1-p['residual'])/p['dep_years']*dep_active
    cumulative_land=np.full(31,land_cost)
    cumulative_nonland=np.cumsum(nonland*build)
    cumulative_finance=np.cumsum(cap_interest)
    book_project=((cumulative_land+cumulative_nonland)*(1-p['sale_share'])-np.cumsum(dep_project))*on
    book_equity=((cumulative_land+cumulative_nonland+cumulative_finance)*(1-p['sale_share'])-np.cumsum(dep_equity))*on
    terminal_gross=np.zeros(31);terminal_fee=np.zeros(31)
    gain_project=np.zeros(31);gain_equity=np.zeros(31)
    if p['terminal_on']:
        terminal_gross[n]=p['terminal_value'];terminal_fee[n]=p['terminal_value']*p['terminal_fee']
        gain_project[n]=terminal_gross[n]-terminal_fee[n]-book_project[n]
        gain_equity[n]=terminal_gross[n]-terminal_fee[n]-book_equity[n]
    terminal=terminal_gross-terminal_fee
    ebit_project=revenue-cash_cost-sale_cost_project-dep_project
    ebit_equity=revenue-cash_cost-sale_cost_equity-dep_equity
    pretax=ebit_equity-op_interest
    tax_project=np.maximum(0,ebit_project+gain_project)*p['income_tax']
    tax_equity=np.maximum(0,pretax+gain_equity)*p['income_tax']
    profit=pretax+gain_equity-tax_equity
    cfads=revenue-cash_cost-np.maximum(0,pretax)*p['income_tax']
    op_project=revenue-cash_cost-np.maximum(0,ebit_project)*p['income_tax']
    land_cash=np.zeros(31);land_cash[0]=land_cost
    nonland_cash=nonland*build
    wc_in=np.zeros(31);wc_out=np.zeros(31)
    if p['build_years']+1<=n:wc_in[int(p['build_years'])+1]=p['working_capital']
    if n>p['build_years']:wc_out[n]=p['working_capital']
    project=revenue-cash_cost-tax_project-land_cash-nonland_cash-wc_in+wc_out+terminal
    net=project+tax_project-tax_equity+draw-interest-principal
    # Independent reserve calculation: largest cumulative future shortfall.
    reserve=np.array([max(0,-min(np.cumsum(net[t+1:n+1]),default=0)) if t<=n else 0 for t in y])
    required=np.zeros(31);dividend=np.zeros(31);cash=np.zeros(31)
    actual=np.zeros(31);actual_div=np.zeros(31);actual_cash=np.zeros(31);gap=np.zeros(31)
    max_gap=0;actual_total=0
    for t in range(n+1):
        available=(cash[t-1] if t else 0)+net[t]
        required[t]=max(0,-available)
        dividend[t]=max(0,available+required[t]-reserve[t])
        cash[t]=available+required[t]-dividend[t]
        before=(actual_cash[t-1] if t else 0)+net[t]-(gap[t-1] if t else 0)
        actual[t]=min(max(0,-before),max(0,p['equity_cap']-actual_total))
        actual_total+=actual[t]
        gap[t]=max(0,-before-actual[t]);max_gap=max(max_gap,gap[t])
        actual_div[t]=max(0,before+actual[t]-reserve[t]) if max_gap<=.01 else 0
        actual_cash[t]=max(0,before+actual[t])-actual_div[t]
    equity=dividend-required
    discount=project/(1+p['project_discount'])**y
    dscr=np.array([cfads[t]/service[t] if t<=n and t>p['build_years'] and service[t]>1e-6 else np.nan for t in y])
    total=development+p['working_capital']
    cum_sale=np.cumsum(sales_plan)*on
    area=dict(land=land,above=above,operable=operable,sale=sold_area,held=held_area,under=p['under_area'])
    inv=dict(land=land_cost,above=above_cost,under=under_cost,infra=p['infra']*factor,pre=p['pre']*factor,
             other=p['other_build']*factor,pre_cont=pre_cont,contingency=pre_cont*p['contingency'],
             nonland=nonland,cap_interest=capitalized,development=development,wc=p['working_capital'],total=total,
             sale_basis_project=sale_basis_project,dep_basis_project=dep_basis_project,held_land=held_land,
             sale_basis_equity=sale_basis_equity,dep_basis_equity=dep_basis_equity,held_invest=held_invest)
    inv_annual=dict(land_cash=land_cash,nonland_cash=nonland_cash,build_cash=land_cash+nonland_cash,wc_in=wc_in,
                    wc_out=wc_out,project_invest_cash=land_cash+nonland_cash+wc_in,
                    finance_cost_cash=cap_interest,total_cash=land_cash+nonland_cash+wc_in+cap_interest)
    plan=dict(build=build,sale=sales_plan,cum_sale=cum_sale,sold=sold_area*sales_plan,
              unsold=sold_area*(1-cum_sale)*on,rent_age=age,occ=occ,occupied=occupied,new=new,
              complete=(cum_sale>=.999999).astype(float)*on)
    inc=dict(sale_ratio=sales_plan,sale_area=sold_area*sales_plan,sale_price=sale_price*on,sales=sale_revenue,
             occ=occ,occupied=occupied,new=new,rent_age=age,rent_price=price,free_loss=free_loss,rent_cash=rent,
             other=other,held_income=held_income,revenue=revenue,terminal_gross=terminal_gross,
             terminal_fee=terminal_fee,terminal_net=terminal)
    costs=dict(sale_expense=sale_expense,sale_tax=sale_tax,fixed=fixed,variable=variable,maintenance=maintenance,
               operating_tax=op_tax,cash_total=cash_cost,sale_cost_project=sale_cost_project,
               sale_cost_equity=sale_cost_equity,dep_project=dep_project,dep_equity=dep_equity,
               book_project=book_project,book_equity=book_equity,
               inventory=((cumulative_land+cumulative_nonland+cumulative_finance)*p['sale_share']-np.cumsum(sale_cost_equity))*on,
               gain_project=gain_project,gain_equity=gain_equity)
    finance=dict(opening=opening,draw=draw,draw_flag=(draw>1e-6).astype(float),principal=principal,
                 interest=interest,closing=closing,cap_interest=cap_interest,op_interest=op_interest,
                 service=service,maturity_balance=closing*(y==maturity)*on)
    cf=dict(revenue=revenue,cash_cost=cash_cost,ebit_project=ebit_project,ebit_equity=ebit_equity,
            op_interest=op_interest,pretax=pretax,tax_project=tax_project,tax_equity=tax_equity,profit=profit,
            cfads=cfads,op_project=op_project,land=land_cash,build=nonland_cash,wc_in=wc_in,wc_out=wc_out,
            terminal=terminal,project=project,draw=draw,interest=interest,principal=principal,
            tax_bridge=tax_project-tax_equity,before_equity=net,reserve=reserve,
            theory_open=np.r_[0,cash[:-1]]*on,required_equity=required,theory_distribution=dividend,
            theory_close=cash,cum_required=np.cumsum(required)*on,equity=equity,
            actual_open=np.r_[0,actual_cash[:-1]]*on,gap_open=np.r_[0,gap[:-1]]*on,
            actual_equity=actual,actual_distribution=actual_div,actual_close=actual_cash,gap=gap,
            max_gap=np.maximum.accumulate(gap)*on,cum_actual=np.cumsum(actual)*on,
            actual_equity_cf=actual_div-actual,cum_project=np.cumsum(project)*on,discounted=discount,
            cum_discounted=np.cumsum(discount)*on,dscr=dscr,
            sale_cash_contribution=sale_revenue-sale_expense-sale_tax,
            held_cash_contribution=held_income-fixed-variable-maintenance-op_tax)
    if n>=p['rent_start']+3:
        stable=int(p['rent_start']+3)
        full_income=held_area*price[stable]*12/10000*p['collection']*(1+p['other_income'])
        denom=full_income*(1-p['variable_cost']-p['operating_tax'])
        be=(fixed[stable]+p['maintenance'])/denom if denom>0 else '无法计算'
        noi=held_income[stable]-fixed[stable]-variable[stable]-maintenance[stable]-op_tax[stable]
        noi_yield=noi/held_invest if held_invest>0 else '无法计算'
        stable_occ=occ[stable]
    else:stable=be=noi_yield=stable_occ='不适用'
    total_in=sum(revenue)+sum(terminal)+sum(wc_out)
    metrics=dict(project_irr=irr(project[:n+1]),equity_irr=irr(equity[:n+1]),
                 project_npv=float(sum(discount)),equity_npv=float(sum(equity/(1+p['equity_discount'])**y)),
                 static_payback=payback(project[:n+1]),dynamic_payback=payback(project[:n+1],p['project_discount']),
                 dscr=float(np.nanmin(dscr)) if p['loan']>0 and any(np.isfinite(dscr)) else '不适用' if p['loan']==0 else '无法计算',
                 required_equity=float(sum(required)),max_gap=float(max(gap)),
                 maturity_balance=closing[int(maturity)] if p['loan']>0 and maturity<=n and maturity<=30 else '无法计算' if p['loan']>0 else 0,
                 total_investment=total,loan=p['loan'],equity_cap=p['equity_cap'],sale_through=sum(sales_plan),
                 sale_complete=next((int(t) for t in y if cum_sale[t]>=.999999),'测算期内未完成'),
                 stable_occ=stable_occ,breakeven=be,noi_yield=noi_yield,
                 terminal_share=sum(terminal)/total_in if total_in else '不适用',stable_year=stable)
    for key,flows in [('project',project),('equity',equity)]:
        signs=np.sign(flows[np.abs(flows)>1e-6]);changes=sum(signs[1:]!=signs[:-1])
        metrics[key+'_sign_warning']='多次变号：IRR可能有歧义' if changes>1 else '单次变号或无变号'
    metrics['equity_label']='理论结果，资金尚未落实' if max(gap)>.01 else '按完整出资计划测算'
    return dict(area=area,investment=inv,investment_annual=inv_annual,plan=plan,income=inc,costs=costs,
                financing=finance,financing_scalar=dict(first_draw=first_draw,maturity=maturity,repay_count=count),
                cashflow=cf,metrics=metrics,params=p)

def status(metric,threshold,kind):
    if metric=='不适用':return '不适用' if kind=='dscr' else '未通过'
    if isinstance(metric,str):return '待判断' if metric=='无法计算' else '未通过'
    ok=(metric>=threshold-1e-8 if kind in ['ge','dscr'] else metric<=threshold+1e-8 if kind=='le'
        else metric<=threshold+.01 if kind=='le_money' else metric>=-.01 if kind=='nonnegative' else abs(metric)<=.01)
    return '通过' if ok else '未通过'

def prepare(book,work):
    work.mkdir(parents=True,exist_ok=True);copy=work/'Excel实际测试.xlsx';shutil.copy2(book,copy)
    manifest=dict(defaults=g.DEFAULTS,input_rows=g.INPUT_ROWS,build_row=g.BUILD_PLAN_ROW,sale_row=g.SALE_PLAN_ROW,
                  build=g.DEFAULT_BUILD_PLAN,sales=g.DEFAULT_SALE_PLAN,cases=[dict(name=n,overrides=o) for n,o in CASES],
                  sheets=g.SHEETS,offset=g.OFFSET,
                  grids={'plan':[2,list(g.PLAN.values())], 'investment_annual':[3,list(g.INV.values())[19:]],
                         'income':[4,list(g.INC.values())], 'costs':[5,list(g.COST.values())],
                         'financing':[6,list(g.FIN.values())[3:]], 'cashflow':[7,list(g.CF.values())]},
                  scalars={'area':[2,list(g.AREA.values())], 'investment':[3,list(g.INV.values())[:19]],
                           'financing_scalar':[6,list(g.FIN.values())[:3]]})
    (work/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8-sig')
    ps=r'''param([string]$WorkDir)
$ErrorActionPreference='Stop'
$m=Get-Content -LiteralPath (Join-Path $WorkDir 'manifest.json') -Raw -Encoding UTF8 | ConvertFrom-Json
$excel=$null;$book=$null;$reports=New-Object System.Collections.Generic.List[object]
function Reset-Inputs {
 foreach($property in $m.defaults.PSObject.Properties) {
  $row=$m.input_rows.($property.Name)
  if($property.Value -is [string]) {
   $book.Worksheets.Item(2).Range('B'+$row).Value2=[string]$property.Value
  } else {
   $book.Worksheets.Item(2).Range('B'+$row).Value2=[double]$property.Value
  }
 }
 for($y=0;$y -le 30;$y++) {
  $book.Worksheets.Item(2).Cells.Item($m.build_row,4+$y).Value2=[double]$m.build[$y]
  $book.Worksheets.Item(2).Cells.Item($m.sale_row,4+$y).Value2=[double]$m.sales[$y]
 }
}
function Snapshot([string]$name) {
 $r=[ordered]@{name=$name;scenarios=@();checks=@();decision=$book.Worksheets.Item(10).Range('A5').Value2;
              breakeven_notice=$book.Worksheets.Item(9).Range('G22').Value2}
 foreach($off in @(0,80)) {
  $s=[ordered]@{}
  foreach($pr in $m.grids.PSObject.Properties) {
   $sheet=$book.Worksheets.Item(1+[int]$pr.Value[0]);$rows=[ordered]@{}
   foreach($row in $pr.Value[1]) {
    $vals=$sheet.Range($sheet.Cells.Item([int]$row+$off,4),$sheet.Cells.Item([int]$row+$off,34)).Value2
    $a=@();for($j=1;$j -le 31;$j++) { $a+=,$vals.GetValue(1,$j) };$rows.([string]$row)=$a
   };$s.($pr.Name)=$rows
  }
  foreach($pr in $m.scalars.PSObject.Properties) {
   $sheet=$book.Worksheets.Item(1+[int]$pr.Value[0]);$rows=[ordered]@{}
   foreach($row in $pr.Value[1]) { $rows.([string]$row)=$sheet.Cells.Item([int]$row+$(if($pr.Name -eq 'area'){0}else{$off}),2).Value2 };$s.($pr.Name)=$rows
  }
  $col=$(if($off -eq 0){2}else{4});$statusCol=$(if($off -eq 0){4}else{6})
  $s.metrics=[ordered]@{};for($row=6;$row -le 28;$row++) { $s.metrics.([string]$row)=$book.Worksheets.Item(9).Cells.Item($row,$col).Value2 }
  $s.status=@();for($row=34;$row -le 43;$row++) {$s.status+=,$book.Worksheets.Item(9).Cells.Item($row,$statusCol).Value2}
  $s.overall=$book.Worksheets.Item(9).Cells.Item(46,$(if($off -eq 0){2}else{5})).Value2
  $s.reasons=$book.Worksheets.Item(9).Cells.Item($(if($off -eq 0){48}else{50}),2).Value2
  $r.scenarios+=,$s
 }
 for($row=35;$row -le 64;$row++) {$r.checks+=,$book.Worksheets.Item(10).Cells.Item($row,4).Value2}
 $r.chartPoints=@();foreach($chart in $book.Worksheets.Item(10).ChartObjects()) { $r.chartPoints+=,$chart.Chart.SeriesCollection(1).Points().Count }
 return $r
}
try {
 $excel=New-Object -ComObject Excel.Application
 $excel.Visible=$false;$excel.DisplayAlerts=$false;$excel.AskToUpdateLinks=$false
 $book=$excel.Workbooks.Open((Join-Path $WorkDir 'Excel实际测试.xlsx'),0,$false)
 $excel.Calculation=-4105
 foreach($case in $m.cases) {
  Reset-Inputs
  foreach($pr in $case.overrides.PSObject.Properties) {
   if($pr.Name -eq 'sale_plan') {for($y=0;$y -le 30;$y++) {$book.Worksheets.Item(2).Cells.Item($m.sale_row,4+$y).Value2=[double]$pr.Value[$y]}}
   else { $book.Worksheets.Item(2).Cells.Item([int]$m.input_rows.($pr.Name),2).Value2=[double]$pr.Value }
  }
  $excel.CalculateFullRebuild()
  if($excel.CalculationState -ne 0) {throw 'Excel recalculation incomplete'}
  $r=Snapshot $case.name;$reports.Add($r)
  Write-Output ('Completed: '+$case.name+' / '+$r.decision)
 }
 Reset-Inputs;$excel.CalculateFullRebuild();$book.Save()
 # Export native Excel charts and the opening view of every worksheet for visual QA.
 $qa=Join-Path $WorkDir 'visual_qa';New-Item -ItemType Directory -Path $qa -Force | Out-Null
 for($i=1;$i -le 10;$i++) {
  $sheet=$book.Worksheets.Item($i);$oldArea=$sheet.PageSetup.PrintArea
  $sheet.PageSetup.PrintArea=$(if($i -eq 1){'A1:G36'}elseif($i -eq 2){'A1:K36'}elseif($i -eq 9){'A1:G29'}elseif($i -eq 10){'A1:G32'}else{'A1:J36'})
  $sheet.PageSetup.Zoom=$false;$sheet.PageSetup.FitToPagesWide=1;$sheet.PageSetup.FitToPagesTall=1
  $sheet.ExportAsFixedFormat(0,(Join-Path $qa ('sheet_'+$i+'.pdf')))
  $sheet.PageSetup.PrintArea=$oldArea;$sheet.PageSetup.FitToPagesTall=$false
 }
 $sheet=$book.Worksheets.Item(10);$oldArea=$sheet.PageSetup.PrintArea
 $sheet.PageSetup.PrintArea='A67:H117';$sheet.PageSetup.FitToPagesWide=1;$sheet.PageSetup.FitToPagesTall=$false
 $sheet.ExportAsFixedFormat(0,(Join-Path $qa 'charts.pdf'))
 $sheet.PageSetup.PrintArea=$oldArea
 $r=Snapshot '恢复默认';$reports.Add($r)
 $book.Save()
 [ordered]@{excel_version=$excel.Version;cases=$reports} | ConvertTo-Json -Depth 15 | Set-Content -LiteralPath (Join-Path $WorkDir 'actual_results.json') -Encoding UTF8
} finally {
 if($null -ne $book) {$book.Close($false);[void][Runtime.InteropServices.Marshal]::ReleaseComObject($book)}
 if($null -ne $excel) {$excel.Quit();[void][Runtime.InteropServices.Marshal]::ReleaseComObject($excel)}
 [GC]::Collect();[GC]::WaitForPendingFinalizers()
}
'''
    (work/'run_excel_tests.ps1').write_text(ps,encoding='utf-8-sig')
    print(f'Prepared {len(CASES)} cases in {work.resolve()}')

def compare(a,b,label):
    if isinstance(b,str):assert a==b,f'{label}: actual {a!r}, expected {b!r}'
    elif np.isscalar(b):
        assert isinstance(a,(int,float)) and math.isclose(float(a),float(b),abs_tol=1e-6,rel_tol=2e-7),f'{label}: actual {a!r}, expected {b!r}'
    else:
        assert len(a)==len(b),label
        for i,(av,bv) in enumerate(zip(a,b)):
            if np.isnan(bv):assert av in [None,''],f'{label}[{i}]: DSCR should be blank'
            else:compare(av,bv,f'{label}[{i}]')

def structural_check(path):
    wb=openpyxl.load_workbook(path,data_only=False);values=openpyxl.load_workbook(path,data_only=True)
    assert wb.sheetnames==g.SHEETS
    assert not wb._external_links
    assert len(wb['决策看板']._charts)==3
    input_sheet=wb['核心参数输入']
    expected=[f'B{g.INPUT_ROWS[k]}' for k,v in g.DEFAULTS.items() if not isinstance(v,str)]
    expected += [f'{g.col(y)}{r}' for r in [g.BUILD_PLAN_ROW,g.SALE_PLAN_ROW] for y in range(31)]
    for address in expected:
        assert any(address in rule for rule in input_sheet.data_validations.dataValidation),f'No validation: {address}'
    errors=[];formula_cells={};graph={}
    for ws in wb:
        assert ws.sheet_state=='visible',ws.title
        for row in ws:
            for c in row:
                if values[ws.title][c.coordinate].data_type=='e':errors.append((ws.title,c.coordinate,values[ws.title][c.coordinate].value))
                if c.data_type=='f':formula_cells[(ws.title,c.coordinate)]=c.value
    assert not errors,f'Excel formula errors: {errors[:20]}'
    for node,formula in formula_cells.items():
        deps=set()
        for token in Tokenizer(formula).items:
            if token.type!='OPERAND' or token.subtype!='RANGE':continue
            ref=token.value
            if '!' in ref:
                sheet,ref=ref.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
                assert sheet in wb.sheetnames,f'Unknown sheet {sheet}'
            else:sheet=node[0]
            if ref in wb.defined_names:continue
            assert '#REF!' not in ref
            if not re.fullmatch(r'\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?',ref):continue
            lo,top,hi,bottom=range_boundaries(ref)
            assert hi<=16384 and bottom<=1048576
            for row in wb[sheet].iter_rows(min_row=top,max_row=bottom,min_col=lo,max_col=hi):
                deps.update((sheet,c.coordinate) for c in row if (sheet,c.coordinate) in formula_cells)
        graph[node]=deps
    # Kahn's algorithm tests the actual formula dependency graph without recursion limits.
    pending={k:len(v) for k,v in graph.items()};reverse={k:[] for k in graph}
    for k,deps in graph.items():
        for d in deps:reverse[d].append(k)
    ready=[k for k,v in pending.items() if v==0];count=0
    while ready:
        k=ready.pop();count+=1
        for d in reverse[k]:
            pending[d]-=1
            if pending[d]==0:ready.append(d)
    assert count==len(graph),f'Circular dependencies: {[k for k,v in pending.items() if v][:20]}'
    return dict(formulas=len(formula_cells),sheets=len(wb.sheetnames),charts=3,formula_errors=0,circular_dependencies=0)

def check(book,work):
    raw=json.loads((work/'actual_results.json').read_text(encoding='utf-8-sig'))
    completed=[];comparisons=0
    assert {n for n,_ in CASES}.issubset({r['name'] for r in raw['cases']}),'Missing scenario tests'
    maplist=[('plan',g.PLAN),('investment_annual',{k:v for k,v in g.INV.items() if v>=28}),
             ('income',g.INC),('costs',g.COST),('financing',{k:v for k,v in g.FIN.items() if v>=10}),('cashflow',g.CF)]
    configs={n:o for n,o in CASES};configs['恢复默认']={}
    for result in raw['cases']:
        name=result['name'];config=configs[name]
        for i,stress in enumerate([False,True]):
            ref=oracle(config,stress);actual=result['scenarios'][i];label=f'{name}/{"压力" if stress else "基准"}'
            for group,mapping in maplist:
                for key,expected in ref[group].items():
                    compare(actual[group][str(mapping[key])],expected,label+'/'+group+'/'+key);comparisons+=len(expected)
            for group,mapping in [('area',g.AREA),('investment',g.INV),('financing_scalar',g.FIN)]:
                for key,expected in ref[group].items():compare(actual[group][str(mapping[key])],expected,label+'/'+key);comparisons+=1
            for key,expected in ref['metrics'].items():compare(actual['metrics'][str(g.METRIC_ROWS[key])],expected,label+'/'+key);comparisons+=1
            desired=[status(ref['metrics'][key],ref['params'].get(threshold,0),kind) for key,_,threshold,kind in g.CONDITIONS]
            assert actual['status']==desired,f'{label}: status mismatch {actual["status"]} vs {desired}'
            for st,(_,title,_,_) in zip(desired,g.CONDITIONS):
                if st=='未通过':assert title+'未达标' in actual['reasons'],label+' reason missing '+title
                if st=='待判断':assert title+'无法计算' in actual['reasons'],label+' unavailable reason missing '+title
            # No sales-area rent, repeated construction cash, phantom equity or post-dividend calls.
            assert max(ref['income']['occupied'])<=ref['area']['held']+1e-6
            assert math.isclose(sum(ref['investment_annual']['nonland_cash']),ref['investment']['nonland'],abs_tol=1e-6)
            assert sum(ref['cashflow']['actual_equity'])<=ref['params']['equity_cap']+.01
            for t in range(31):
                if ref['cashflow']['theory_distribution'][t]>.01:
                    assert sum(ref['cashflow']['required_equity'][t+1:])<=.01
        valid=name not in ['期限不覆盖贷款','贷款到期超出预留期']
        if valid:assert all(x=='通过' for x in result['checks']),f'{name}: failed checks {result["checks"]}'
        else:assert '未通过' in result['checks'] and result['decision'].startswith('待判断'),name
        if name=='资本金额度不足':assert result['scenarios'][0]['metrics']['14']>0 and result['decision']=='未通过'
        if name=='无正现金流IRR无解':assert result['decision'].startswith('待判断')
        if name=='模拟全部条件通过':assert result['decision'].startswith('通过'),result['decision']
        if name=='基准通过但要求压力通过':
            assert result['scenarios'][0]['overall']=='通过' and result['scenarios'][1]['overall']=='未通过'
            assert result['decision']=='未通过'
        if name=='保本出租率超过100%':
            assert result['scenarios'][0]['metrics']['22']>1
            assert '保本出租率超过100%' in result['breakeven_notice']
        assert result['chartPoints']==[int(ref['params']['horizon'])+1]*3,f'{name}: charts did not resize {result["chartPoints"]}'
        completed.append(dict(name=name,decision=result['decision'],metrics=result['scenarios'][0]['metrics'],checks_passed=all(x=='通过' for x in result['checks'])))
    # Excel saved workbook must restore the author's defaults, including every yearly plan.
    final=work/'Excel实际测试.xlsx'
    cached=openpyxl.load_workbook(final,data_only=True)
    for k,v in g.DEFAULTS.items():compare(cached[g.SHEETS[1]][f'B{g.INPUT_ROWS[k]}'].value,v,'restored '+k)
    for row,plan in [(g.BUILD_PLAN_ROW,g.DEFAULT_BUILD_PLAN),(g.SALE_PLAN_ROW,g.DEFAULT_SALE_PLAN)]:
        for year,value in enumerate(plan):compare(cached[g.SHEETS[1]].cell(row,4+year).value,value,'restored annual plan')
    structure=structural_check(final)
    summary=dict(excel_version=raw['excel_version'],independent_value_comparisons=comparisons,structure=structure,cases=completed,
                 baseline=raw['cases'][0],default_restored=True)
    (work/'validation_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    # Never overwrite a user's workbook. Promotion of a verified default uses a new path.
    verified=g.unique_path(book.parent,book.name)
    shutil.copy2(final,verified)
    print('Verified workbook:',verified.resolve())
    print(json.dumps(dict(comparisons=comparisons,structure=structure,cases=len(completed),default_restored=True),ensure_ascii=False))

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode',choices=['prepare','check','oracle'])
    parser.add_argument('--book',type=Path,default=Path(__file__).parent/'产业园投资测算.xlsx')
    parser.add_argument('--work-dir',type=Path,default=Path(__file__).parent/'验证临时文件')
    args=parser.parse_args()
    if args.mode=='prepare':prepare(args.book,args.work_dir)
    elif args.mode=='check':check(args.book,args.work_dir)
    else:print(json.dumps(oracle()['metrics'],ensure_ascii=False,indent=2))
if __name__=='__main__':main()
