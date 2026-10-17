from pathlib import Path
import json,zipfile,hashlib,xml.etree.ElementTree as E
import openpyxl
root=Path(__file__).parent;path=root.parent/'产业园投资测算_V3_决策测算.xlsx'
w=openpyxl.load_workbook(path,data_only=False);v=openpyxl.load_workbook(path,data_only=True);s=v['敏感性与边界']
z=zipfile.ZipFile(path);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
tables=[];chart_refs=[]
for name in z.namelist():
 if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
  for x in E.fromstring(z.read(name)).findall('.//s:f',ns):
   if x.get('t')=='dataTable':tables.append({'part':name,**x.attrib})
 if name.startswith('xl/charts/chart') and name.endswith('.xml'):
  chart_refs.append({'part':name,'refs':[x.text for x in E.fromstring(z.read(name)).findall('.//c:f',ns)]})
m=json.loads((root/'manifest.json').read_text(encoding='utf-8'))
source_hash=hashlib.sha256(Path(m['source']['source']).read_bytes()).hexdigest()
result={'file':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'source_unchanged':source_hash==m['source']['sha256'],
 'bounds':[{'name':s.cell(r,1).value,'value':s.cell(r,2).value,'note':s.cell(r,4).value} for r in range(73,82)],
 'boundary_status':s['B66'].value,'refreshed_at':s['B67'].value,'freeze':[(s.title,str(s.freeze_panes),s.protection.sheet) for s in w],
 'input_unlocked':not w['核心参数输入']['B36'].protection.locked,'formula_locked':w['核心参数输入']['B15'].protection.locked,
 'named_ranges':{k:x.attr_text for k,x in w.defined_names.items()},'native_data_tables':tables,'charts':chart_refs,
 'external_parts':[n for n in z.namelist() if 'externalLink' in n or 'vbaProject' in n]}
(root/'final_review.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2,default=str))
