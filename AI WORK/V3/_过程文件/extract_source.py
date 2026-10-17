from pathlib import Path
import json,hashlib
import openpyxl
r=Path(__file__).resolve().parents[2]
p=r/'V2/产业园投资测算_V2_样式调整.xlsx'
w=openpyxl.load_workbook(p,data_only=False)
v=openpyxl.load_workbook(p,data_only=True)
s=w['核心参数输入']
d={'source':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'parameters':[],
   'build':[v[s.title].cell(88,c).value for c in range(4,35)],'sale':[v[s.title].cell(89,c).value for c in range(4,35)]}
for row in range(7,83):
 if s.cell(row,2).value is not None:
  d['parameters'].append({'row':row,'name':s.cell(row,1).value,'value':v[s.title].cell(row,2).value,'formula':s.cell(row,2).value if s.cell(row,2).data_type=='f' else None,'unit':s.cell(row,3).value,'note':s.cell(row,4).value})
(Path(__file__).parent/'source.json').write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'sha256':d['sha256'],'parameters':len(d['parameters']),'sheets':w.sheetnames},ensure_ascii=False))
