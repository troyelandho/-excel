from pathlib import Path
import json,re,math,sys
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries,get_column_letter
root=Path(__file__).parent
path=root/'候选_V3.xlsx'
if len(sys.argv)>1:path=Path(sys.argv[1])
w=openpyxl.load_workbook(path,data_only=False)
m=json.loads((root/'manifest.json').read_text(encoding='utf-8'))
formulas={(s.title,c.coordinate):c.value for s in w for row in s for c in row if c.data_type=='f' and isinstance(c.value,str)}
graph={};invalid=[]
for key,formula in formulas.items():
 deps=set()
 try:tokens=Tokenizer(formula).items
 except Exception as e:invalid.append((key,str(e),formula));continue
 for token in tokens:
  if token.type!='OPERAND' or token.subtype!='RANGE':continue
  ref=token.value
  if '!' in ref:s,r=ref.rsplit('!',1);s=s.strip("'")
  else:s,r=key[0],ref
  try:
   c1,r1,c2,r2=range_boundaries(r)
   for row in range(r1,r2+1):
    for col in range(c1,c2+1):
     cell=(s,get_column_letter(col)+str(row))
     if cell in formulas:deps.add(cell)
  except Exception as e:invalid.append((key,ref,str(e)))
 graph[key]=deps
state={};cycles=[]
def visit(key,path):
 if state.get(key)==1:cycles.append(path[path.index(key):]+[key]);return
 if state.get(key)==2:return
 state[key]=1
 for dep in graph.get(key,[]):visit(dep,path+[key])
 state[key]=2
for key in graph:visit(key,[])
source_diffs=[]
for p in m['source']['parameters']:
 if p['formula']:continue
 got=w['核心参数输入'][f'B{p["row"]}'].value
 if got!=p['value']:source_diffs.append((p['row'],p['value'],got))
report={'formula_count':len(formulas),'cycle_count':len(cycles),'cycles':cycles[:5],'invalid':invalid[:15],'source_input_differences':source_diffs,
        'sheets':w.sheetnames,'chart_count':sum(len(s._charts) for s in w),'data_table_formulas':sum('TABLE(' in x for x in formulas.values())}
(root/'static_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
