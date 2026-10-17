import json,math
from pathlib import Path
from openpyxl.utils.cell import get_column_letter

root=Path(__file__).parent
def irr(flows):
 lo,hi=-0.9,10
 def f(r):return sum(v/(1+r)**t for t,v in enumerate(flows))
 for _ in range(150):
  mid=(lo+hi)/2
  if f(mid)>0:lo=mid
  else:hi=mid
 return (lo+hi)/2
d=json.loads((root/'workbook.json').read_text(encoding='utf-8'))
ss={s['name']:s['cells'] for s in d['sheets']}
p=lambda r:ss['核心参数输入']['B'+str(r)]['v']
N=int(p(9));K=int(p(16)); T=range(N+1)
landarea=p(10); ground=landarea*p(11);usable=ground*p(12); sold=usable*p(14);held=usable*(1-p(14))
report={};fail=[]
def compare(sheet,row,values,offset=0):
 for t,v in enumerate(values):
  addr=get_column_letter(t+4)+str(row+offset)
  actual=ss[sheet][addr]['v']
  if actual is None and v is None:continue
  if isinstance(v,(float,int)) and isinstance(actual,(float,int)) and math.isclose(v,actual,abs_tol=1e-7,rel_tol=1e-9):continue
  if v==actual:continue
  fail.append((sheet,addr,v,actual))
def check_metric(row,val,pressure=False):
 addr=('D' if pressure else 'B')+str(row); actual=ss['指标与情景测算'][addr]['v']
 if isinstance(val,(float,int)) and isinstance(actual,(float,int)) and math.isclose(val,actual,abs_tol=1e-7,rel_tol=1e-9):return
 if val!=actual:fail.append(('指标与情景测算',addr,val,actual))
for stress in [False,True]:
 o=80 if stress else 0; f=1+p(81) if stress else 1
 land=landarea*p(21)/10000
 nonland=(ground*p(22)/10000+p(13)*p(23)/10000+p(24)+p(25)+p(26))*(1+p(27))*f
 cp=[ss['核心参数输入'][get_column_letter(t+4)+'88']['v'] if 1<=t<=K else 0 for t in T]
 sp=[ss['核心参数输入'][get_column_letter(t+4)+'89']['v'] if t>=p(17) else 0 for t in T]
 draw=[p(57)*x for x in cp]
 first=next((t for t,x in enumerate(draw) if x>1e-6),0)
 maturity=first+p(59)-1 if p(57)>0 else 0
 repayment_years=max(1,maturity-p(60)+1)
 opening=[]; principal=[]; interest=[];closing=[]; ci=[]; oi=[]
 for t in T:
  op=closing[-1] if t else 0
  pr=min(op+draw[t],p(57)/repayment_years) if p(57)>0 and p(60)<=t<=maturity else 0
  it=max(0,op+draw[t]/2-pr/2)*p(58)
  opening.append(op);principal.append(pr);interest.append(it);closing.append(max(0,op+draw[t]-pr));ci.append(it if t<=K else 0);oi.append(it if t>K else 0)
 capinterest=sum(ci)
 rental=[]; extra=[];sales=[];occ=[];rent=[];maintenance=[];fixed=[];variable=[];optax=[];salefee=[];saletax=[];cashcost=[];depp=[];depe=[];scostp=[];scoste=[];bookp=[];booke=[];inventory=[];invest=[];wkout=[];wkin=[];disposal=[];projtax=[];equitytax=[];cfads=[];pcf=[];eqpre=[];netprofit=[];ecum=0;accd_p=0;accd_e=0
 for t in T:
  rental_age=t-p(18)+1 if t>=p(18) else 0
  oc=(p(38) if rental_age==1 else p(39) if rental_age==2 else p(40) if rental_age==3 else p(41) if rental_age>3 else 0)
  oc=max(0,oc-p(80)/100) if stress and rental_age>0 else oc
  occ.append(oc)
  monthly=p(36)*(1+p(37))**(rental_age-1)*(1-p(79) if stress else 1) if rental_age else 0
  rent.append(monthly)
  new_area=held*max(0,oc-(occ[t-1] if t else 0))
  r=(held*oc*monthly*12-new_area*monthly*p(42))/10000*p(43)
  rental.append(r);extra.append(r*p(44))
  sale=sold*sp[t]*p(31)*(1-p(78) if stress else 1)/10000;sales.append(sale)
  fixed.append(p(47)*(1+p(48))**(t-K-1) if t>K else 0)
  maintenance.append(p(50) if t>K else 0)
  variable.append((r+extra[-1])*p(49));optax.append((r+extra[-1])*p(51));salefee.append(sale*p(32));saletax.append(sale*p(33))
  cash=sum([fixed[-1],maintenance[-1],variable[-1],optax[-1],salefee[-1],saletax[-1]]);cashcost.append(cash)
  dp=nonland*(1-p(14))*(1-p(54))/p(53) if 0<t-K<=p(53) else 0
  de=(nonland+capinterest)*(1-p(14))*(1-p(54))/p(53) if 0<t-K<=p(53) else 0
  depp.append(dp);depe.append(de);accd_p+=dp;accd_e+=de
  scp=(land+nonland)*p(14)*sp[t];sce=(land+nonland+capinterest)*p(14)*sp[t]
  scostp.append(scp);scoste.append(sce)
  paid=land+nonland*sum(cp[:t+1]); paid_e=paid+sum(ci[:t+1])
  bookp.append(paid*(1-p(14))-accd_p);booke.append(paid_e*(1-p(14))-accd_e)
  inventory.append(paid_e*p(14)-sum(scoste))
  invest.append((land if t==0 else 0)+nonland*cp[t])
  wkout.append(p(28) if t==K+1 else 0);wkin.append(p(28) if t==N and N>K else 0)
  dv=p(65)*(1-p(66)) if p(64)==1 and t==N else 0;disposal.append(dv)
  gainp=dv-bookp[-1] if p(64)==1 and t==N else 0
  gaine=dv-booke[-1] if p(64)==1 and t==N else 0
  total=sale+r+extra[-1]
  ebitp=total-cash-scp-dp;ebite=total-cash-sce-de;pretax=ebite-oi[t]
  tp=max(0,ebitp+gainp)*p(52);te=max(0,pretax+gaine)*p(52)
  projtax.append(tp);equitytax.append(te);netprofit.append(pretax+gaine-te)
  cfads.append(total-cash-max(0,pretax)*p(52))
  pc=total-cash-tp-invest[-1]-wkout[-1]+wkin[-1]+dv;pcf.append(pc)
  eqpre.append(total-cash-te-invest[-1]-wkout[-1]+wkin[-1]+dv+draw[t]-interest[t]-principal[t])
 reserve=[0]*(N+1)
 for t in range(N-1,-1,-1):reserve[t]=max(0,reserve[t+1]-eqpre[t+1])
 capital=[];distribution=[];cashbal=[];capped=[];capdist=[];capcash=[];gap=[];evergap=[];lastbal=0;lastgap=0;cumcap=0;historygap=0
 for t in T:
  op=cashbal[-1] if t else 0
  eq=max(0,-op-eqpre[t]);capital.append(eq)
  dist=max(0,op+eqpre[t]+eq-reserve[t]);distribution.append(dist);cashbal.append(op+eqpre[t]+eq-dist)
  need=max(0,-(lastbal+eqpre[t]-lastgap));ce=min(need,max(0,p(61)-cumcap));capped.append(ce);cumcap+=ce
  old=lastbal+eqpre[t]-lastgap+ce;g=max(0,-old);gap.append(g);historygap=max(historygap,g);evergap.append(historygap)
  cd=0 if historygap>.01 else max(0,old-reserve[t]);capdist.append(cd);lastbal=max(0,old)-cd;capcash.append(lastbal);lastgap=g
 ecf=[a-b for a,b in zip(distribution,capital)]
 checks=[('融资与还本付息',10,opening),('融资与还本付息',11,draw),('融资与还本付息',13,principal),('融资与还本付息',14,interest),('融资与还本付息',15,closing),('销售与出租收入',13,sales),('销售与出租收入',20,rental),('运营成本与税费',16,cashcost),('运营成本与税费',19,depp),('运营成本与税费',20,depe),('运营成本与税费',21,bookp),('运营成本与税费',22,booke),('运营成本与税费',23,inventory),('利润与现金流',16,projtax),('利润与现金流',17,equitytax),('利润与现金流',18,netprofit),('利润与现金流',19,cfads),('利润与现金流',26,pcf),('利润与现金流',31,eqpre),('利润与现金流',32,reserve),('利润与现金流',34,capital),('利润与现金流',35,distribution),('利润与现金流',36,cashbal),('利润与现金流',38,ecf),('利润与现金流',42,capped),('利润与现金流',43,capdist),('利润与现金流',44,capcash),('利润与现金流',45,gap)]
 for sheet,row,vals in checks:compare(sheet,row,vals,o)
 npv_p=sum(x/(1+p(71))**t for t,x in enumerate(pcf));npv_e=sum(x/(1+p(72))**t for t,x in enumerate(ecf))
 dscr=[cfads[t]/(principal[t]+interest[t]) for t in T if t>K and principal[t]+interest[t]>1e-6]
 metrics={6:irr(pcf),7:irr(ecf),8:npv_p,9:npv_e,12:min(dscr) if dscr else '不适用',13:sum(capital),14:max(gap),16:land+nonland+capinterest+p(28)}
 for row,v in metrics.items():check_metric(row,v,stress)
 report['stress' if stress else 'base']={
  'metrics':metrics,'independent_checked_cells':len(checks)*(N+1)+len(metrics),
  'equity_by_year':capital,'draw_by_year':draw,'project_cf':pcf,'equity_cf':ecf,
  'max_cash_occupation_before_finance':max(-sum(pcf[:i+1]) for i in T),
  'max_unfunded_gap_year':max(T,key=lambda t:gap[t]),'reserve_peak':max(cashbal),'reserve_peak_year':max(T,key=lambda t:cashbal[t]),
  'dscr_year':min((t for t in T if t>K and principal[t]+interest[t]>1e-6),key=lambda t:cfads[t]/(principal[t]+interest[t])),
  'cum_net_profit':sum(netprofit),'sale_net_before_income_tax':sum(a-b-c for a,b,c in zip(sales,salefee,saletax)),
  'held_noi_before_income_tax':sum(a+b-c-e-f-g for a,b,c,e,f,g in zip(rental,extra,fixed,maintenance,variable,optax)),
  'end_held_book_value':booke[-1], 'end_unsold_inventory':inventory[-1]}
report['differences']=fail
(root/'independent_results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
